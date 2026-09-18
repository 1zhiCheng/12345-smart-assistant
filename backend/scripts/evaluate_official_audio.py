"""比赛真实录音端到端评测：本地 ASR/角色分离 → 工单字段 → 分类与办理单位。

Excel 只提供整理后的工单，不是逐字转写或说话人时间戳标注，因此本脚本不会伪造
CER/WER/DER。公开报告不写入原始编号、音频路径、转写文本或群众诉求正文。
"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from statistics import mean, median
from typing import Any

from openpyxl import load_workbook

from app.config import get_settings
from app.domain.wuhu import CATEGORY_TO_DEPT, seed_wuhu_departments
from app.harness.agents.dept_router import DeptRouter
from app.intake.asr import AsrService
from app.intake.llm_service import LLMIntakeService
from app.intake.service import IntakeService
from app.llm.deepseek import DeepSeekClient
from app.storage.store import MemoryStore
from scripts.evaluate_intake import FIELDS, field_matches
from scripts.evaluate_synthetic_audio import (
    FIELD_THRESHOLDS,
    ngram_scores,
    rounded,
    text_similarity,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATASET = PROJECT_ROOT / "12345赛题数据集-信件类别示例工单及录音" / "信件类别示例工单及录音"
DEFAULT_LABELS = PROJECT_ROOT / "backend" / "evaluation" / "intake_cases.official.deidentified.json"
DEFAULT_OUTPUT = PROJECT_ROOT / "docs" / "competition" / "wuhu_official_audio_e2e_evaluation.json"
SHANGHAI = timezone(timedelta(hours=8), name="Asia/Shanghai")
AUDIO_SUFFIXES = {".mp3", ".wav", ".m4a", ".ogg", ".webm", ".mp4"}
REQUIRED_EXCEL_FIELDS = ("信件编号", "受理时间", "办理单位", "区域")
FORBIDDEN_PUBLIC_KEYS = {
    "audio_file", "audio_path", "source_id", "raw_transcript", "transcript", "text",
    "content", "handling_units", "reply",
}


def _received_at(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value or "").strip().replace("/", "-")
        parsed = None
        for fmt in (None, "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                parsed = datetime.fromisoformat(text) if fmt is None else datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            raise ValueError(f"无法解析受理时间：{text[:32]}")
    return parsed.replace(tzinfo=SHANGHAI) if parsed.tzinfo is None else parsed.astimezone(SHANGHAI)


def _time_key(value: Any) -> str:
    return _received_at(value).isoformat(timespec="seconds")


def _case_number(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _find_audio(category_dir: Path, source_id: str) -> Path:
    candidates = [
        path for path in category_dir.iterdir()
        if path.is_file() and path.suffix.lower() in AUDIO_SUFFIXES and path.stem == source_id
    ]
    if len(candidates) != 1:
        raise ValueError(f"音频匹配应为 1 条，实际 {len(candidates)} 条（类别：{category_dir.name}）")
    return candidates[0]


def read_excel_records(dataset: Path) -> list[dict[str, Any]]:
    """只读取匹配与评测必需字段；不把原始主题、内容或答复写入内存记录。"""
    records: list[dict[str, Any]] = []
    for workbook_path in sorted(dataset.glob("*/*.xlsx")):
        category = workbook_path.parent.name
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                header_index = next(
                    (index for index, row in enumerate(rows) if "信件编号" in {str(cell or "").strip() for cell in row}),
                    None,
                )
                if header_index is None:
                    continue
                headers = [str(value or "").strip() for value in rows[header_index]]
                missing = [field for field in REQUIRED_EXCEL_FIELDS if field not in headers]
                if missing:
                    raise ValueError(f"{workbook_path.name}/{sheet.title} 缺少列：{','.join(missing)}")
                positions = {name: headers.index(name) for name in REQUIRED_EXCEL_FIELDS}
                for row in rows[header_index + 1:]:
                    source_id = _case_number(row[positions["信件编号"]])
                    if not source_id:
                        continue
                    records.append({
                        "category": category,
                        "received_at": _time_key(row[positions["受理时间"]]),
                        "handling_units": [
                            part.strip() for part in re.split(r"[、,，;/；]+", str(row[positions["办理单位"]] or ""))
                            if part.strip()
                        ],
                        "region": str(row[positions["区域"]] or "").strip(),
                        "audio_path": _find_audio(workbook_path.parent, source_id),
                    })
        finally:
            workbook.close()
    return records


def align_cases(dataset: Path, labels_path: Path) -> list[dict[str, Any]]:
    labels = json.loads(labels_path.read_text(encoding="utf-8"))
    label_index: dict[tuple[str, str], dict[str, Any]] = {}
    for case in labels:
        key = (case["official_label"]["category"], _time_key(case["received_at"]))
        if key in label_index:
            raise ValueError(f"脱敏标签存在重复匹配键：{key}")
        label_index[key] = case

    aligned: list[dict[str, Any]] = []
    for record in read_excel_records(dataset):
        key = (record["category"], record["received_at"])
        case = label_index.pop(key, None)
        if case is None:
            raise ValueError(f"Excel 记录没有脱敏标签：{key}")
        if record["region"] != case["official_label"].get("region", ""):
            raise ValueError(f"区域标签不一致：{case['id']}")
        aligned.append({**record, "label": case})
    if label_index:
        raise ValueError(f"仍有 {len(label_index)} 条脱敏标签未匹配到 Excel/音频")
    return aligned


def _unit_aliases(value: str) -> set[str]:
    normalized = re.sub(r"[\s（）()]+", "", value)
    aliases = {normalized}
    for prefix in ("芜湖市", "市"):
        if normalized.startswith(prefix) and len(normalized) > len(prefix):
            aliases.add(normalized[len(prefix):])
    return {alias for alias in aliases if len(alias) >= 3}


def handling_unit_hit(expected: list[str], predicted: list[str]) -> bool:
    expected_aliases = set().union(*(_unit_aliases(unit) for unit in expected)) if expected else set()
    predicted_aliases = set().union(*(_unit_aliases(unit) for unit in predicted)) if predicted else set()
    return any(left in right or right in left for left in expected_aliases for right in predicted_aliases)


def _public_error(exc: Exception) -> str:
    if "本地语音模型不存在" in str(exc):
        return "local_asr_model_missing"
    if "本地语音转写失败" in str(exc):
        return "local_asr_failed"
    return type(exc).__name__


def assert_public_payload(payload: Any) -> None:
    if isinstance(payload, dict):
        blocked = FORBIDDEN_PUBLIC_KEYS.intersection(payload)
        if blocked:
            raise ValueError(f"公开报告包含受限字段：{sorted(blocked)}")
        for value in payload.values():
            assert_public_payload(value)
    elif isinstance(payload, list):
        for value in payload:
            assert_public_payload(value)


async def evaluate_case(
    item: dict[str, Any], asr: AsrService, intake: IntakeService | LLMIntakeService,
    router: DeptRouter, intake_mode: str,
) -> dict[str, Any]:
    started = time.perf_counter()
    case = item["label"]
    audio_path: Path = item["audio_path"]
    public = {
        "id": case["id"],
        "category": item["category"],
        "status": "error",
    }
    try:
        transcript = await asr.transcribe(audio_path.read_bytes(), audio_path.name, "audio/mpeg")
        dialogue = (
            IntakeService.format_diarized_segments(transcript.segments)
            if transcript.diarization_mode == "acoustic"
            else IntakeService.format_audio_dialogue(transcript.text)
        )
        received_at = item["received_at"]
        kwargs = {
            "source_type": "audio",
            "audio_file_name": "official-audio",
            "received_at": received_at,
            "source_channel": "比赛真实热线录音（本地评测）",
            "raw_transcript": transcript.text,
        }
        order = await intake.analyze(dialogue.formatted_text, **kwargs) if intake_mode == "production" else intake.analyze(dialogue.formatted_text, **kwargs)
        query = " ".join(filter(None, [order.title, order.elements.location, order.elements.event, order.elements.request]))
        route = await router.route(query)
        dept_ids = route.get("dept_ids", [])
        dept_names = route.get("dept_names", [])
        expected_dept = CATEGORY_TO_DEPT[item["category"]]
        expected = case["expected"]

        strict_correct = [field for field in FIELDS if expected.get(field) and field_matches(getattr(order.elements, field), expected[field])]
        expected_fields = [field for field in FIELDS if expected.get(field)]
        predicted_fields = [field for field in FIELDS if getattr(order.elements, field)]
        unsupported = [field for field in predicted_fields if field in expected_fields and field not in strict_correct]
        similarities = {
            field: text_similarity("。".join(expected[field]), getattr(order.elements, field))
            for field in ("location", "event", "request")
        }
        passes = {field: similarities[field] >= FIELD_THRESHOLDS[field] for field in similarities}
        citizen_text = "。".join(turn.text for turn in dialogue.turns if turn.role == "citizen")
        support_source = citizen_text or transcript.text
        support_precision, _, _ = ngram_scores(
            support_source,
            f"{order.elements.event}。{order.elements.request}",
        )
        _, official_recall, official_f1 = ngram_scores(case["text"], transcript.text)
        roles = Counter(turn.role for turn in dialogue.turns)
        generation = order.generation.model_dump()
        duration = max((segment.end for segment in transcript.segments), default=0.0)
        public.update({
            "status": "ok",
            "runtime_seconds": rounded(time.perf_counter() - started),
            "recognized_speech_span_seconds": rounded(duration),
            "asr": {
                "provider": transcript.provider,
                "model": transcript.model,
                "recognized_chars": len(re.sub(r"\s+", "", transcript.text)),
                "segment_count": len(transcript.segments),
                "diarization_mode": transcript.diarization_mode,
                "diarization_confidence": rounded(transcript.diarization_confidence),
                "diarization_reason_code": "accepted" if transcript.diarization_mode == "acoustic" else "not_accepted",
            },
            "roles": {
                "operator_turns": roles.get("operator", 0),
                "citizen_turns": roles.get("citizen", 0),
                "unknown_turns": roles.get("unknown", 0),
                "both_roles_present": {"operator", "citizen"}.issubset(roles),
                "format_mode": dialogue.mode,
            },
            "workorder": {
                "generation_mode": generation["mode"],
                "fallback_reason": generation["fallback_reason"],
                "field_passes": {f"{name}_pass": value for name, value in passes.items()},
                "missing_field_names": list(order.missing_fields),
            },
            "routing": {
                "predicted_department_ids": dept_ids,
                "expected_department_id": expected_dept,
                "matched_by": route.get("matched_by", ""),
            },
            "metrics": {
                "official_content_bigram_recall": rounded(official_recall),
                "official_content_bigram_f1": rounded(official_f1),
                "acoustic_diarization_accepted": float(transcript.diarization_mode == "acoustic"),
                "both_roles_present": float({"operator", "citizen"}.issubset(roles)),
                "field_accuracy": rounded(len(strict_correct) / len(expected_fields) if expected_fields else 1.0),
                "information_completeness": rounded(sum(bool(getattr(order.elements, field)) for field in expected_fields) / len(expected_fields) if expected_fields else 1.0),
                "fact_fidelity": rounded(len(strict_correct) / (len(strict_correct) + len(unsupported)) if strict_correct or unsupported else 1.0),
                "location_similarity": rounded(similarities["location"]),
                "event_similarity": rounded(similarities["event"]),
                "request_similarity": rounded(similarities["request"]),
                "field_pass_rate": rounded(mean(passes.values())),
                "workorder_support_precision": rounded(support_precision),
                "region_accuracy": float(case["official_label"]["region"] == order.region or case["official_label"]["region"] in order.elements.location),
                "category_routing_top1": float(bool(dept_ids) and dept_ids[0] == expected_dept),
                "category_routing_topk": float(expected_dept in dept_ids),
                "official_handling_unit_topk": float(handling_unit_hit(item["handling_units"], dept_names)),
            },
        })
    except Exception as exc:  # noqa: BLE001
        public.update({
            "runtime_seconds": rounded(time.perf_counter() - started),
            "error_code": _public_error(exc),
        })
    assert_public_payload(public)
    return public


def _percentile(values: list[float], percentile: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = (len(ordered) - 1) * percentile
    low, high = int(index), min(int(index) + 1, len(ordered) - 1)
    return ordered[low] * (high - index) + ordered[high] * (index - low)


def build_summary(cases: list[dict[str, Any]], expected_total: int) -> dict[str, Any]:
    ok = [case for case in cases if case["status"] == "ok"]
    metric = lambda name: [float(case["metrics"][name]) for case in ok]  # noqa: E731
    by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for case in ok:
        by_category[case["category"]].append(case)
    return {
        "expected_cases": expected_total,
        "evaluated_cases": len(cases),
        "successful_cases": len(ok),
        "failed_cases": len(cases) - len(ok),
        "failed_ids": [case["id"] for case in cases if case["status"] != "ok"],
        "recognized_speech_span_minutes": rounded(sum(case.get("recognized_speech_span_seconds", 0) for case in ok) / 60),
        "runtime_minutes": rounded(sum(case.get("runtime_seconds", 0) for case in cases) / 60),
        "asr": {
            "success_rate": rounded(len(ok) / len(cases)) if cases else 0.0,
            "official_content_bigram_recall_mean": rounded(mean(metric("official_content_bigram_recall"))) if ok else None,
            "official_content_bigram_f1_mean": rounded(mean(metric("official_content_bigram_f1"))) if ok else None,
            "recognized_chars_median": rounded(median(case["asr"]["recognized_chars"] for case in ok)) if ok else None,
        },
        "speaker_roles": {
            "acoustic_diarization_accept_rate": rounded(mean(metric("acoustic_diarization_accepted"))) if ok else None,
            "both_roles_present_rate": rounded(mean(metric("both_roles_present"))) if ok else None,
            "diarization_confidence_mean": rounded(mean(case["asr"]["diarization_confidence"] for case in ok)) if ok else None,
        },
        "workorder": {
            name: rounded(mean(metric(name))) if ok else None for name in (
                "field_accuracy", "information_completeness", "fact_fidelity", "location_similarity",
                "event_similarity", "request_similarity", "field_pass_rate", "workorder_support_precision",
                "region_accuracy",
            )
        },
        "routing": {
            "category_top1_accuracy": rounded(mean(metric("category_routing_top1"))) if ok else None,
            "category_topk_accuracy": rounded(mean(metric("category_routing_topk"))) if ok else None,
            "official_handling_unit_topk_accuracy": rounded(mean(metric("official_handling_unit_topk"))) if ok else None,
            "matched_by": dict(Counter(case["routing"]["matched_by"] for case in ok)),
        },
        "generation_mode_counts": dict(Counter(case["workorder"]["generation_mode"] for case in ok)),
        "runtime_seconds_p95": rounded(_percentile([case["runtime_seconds"] for case in ok], 0.95)) if ok else None,
        "by_category": {
            category: {
                "cases": len(rows),
                "field_accuracy": rounded(mean(row["metrics"]["field_accuracy"] for row in rows)),
                "category_top1_accuracy": rounded(mean(row["metrics"]["category_routing_top1"] for row in rows)),
                "category_topk_accuracy": rounded(mean(row["metrics"]["category_routing_topk"] for row in rows)),
            }
            for category, rows in sorted(by_category.items())
        },
    }


def save_report(output: Path, metadata: dict[str, Any], cases: list[dict[str, Any]], total: int) -> None:
    report = {"metadata": metadata, "summary": build_summary(cases, total), "cases": cases}
    assert_public_payload(report)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


async def run(args: argparse.Namespace) -> dict[str, Any]:
    aligned = align_cases(args.dataset.resolve(), args.labels.resolve())
    if args.validate_only:
        return {
            "excel_records": len(aligned),
            "audio_files": len(aligned),
            "deidentified_labels": len(aligned),
            "categories": dict(sorted(Counter(item["category"] for item in aligned).items())),
            "all_aligned": True,
        }
    if args.limit:
        aligned = aligned[:args.limit]
    if args.intake_mode == "production" and not args.allow_external_llm:
        raise SystemExit("production 模式可能外发脱敏诉求，必须显式添加 --allow-external-llm")

    settings = get_settings().model_copy(update={
        "storage_mode": "memory",
        "vector_backend": "memory",
        "pi_agent_enabled": False,
        "asr_provider": "faster_whisper",
        "asr_local_device": args.device,
        "asr_local_compute_type": args.compute_type,
        **({"asr_local_model_path": str(args.model_path.resolve())} if args.model_path else {}),
        **({"deepseek_api_key": "", "intake_llm_enabled": False} if args.intake_mode == "rules" else {"intake_llm_enabled": True}),
    })
    store = MemoryStore()
    await seed_wuhu_departments(store)
    llm = DeepSeekClient(settings)
    router = DeptRouter(llm, store)
    rules = IntakeService()
    intake: IntakeService | LLMIntakeService = (
        LLMIntakeService(llm, settings, rules) if args.intake_mode == "production" else rules
    )
    asr = AsrService(settings)
    metadata = {
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "dataset_type": "official_real_audio_with_deidentified_excel_labels",
        "intake_mode": args.intake_mode,
        "asr_provider": "faster_whisper",
        "field_similarity_thresholds": FIELD_THRESHOLDS,
        "privacy": "public_safe_no_original_ids_paths_transcripts_or_appeal_text",
        "notes": [
            "Excel 内容是整理后的工单，不是逐字转写，因此不计算 CER/WER。",
            "Excel 没有说话人时间戳，因此不计算 DER；角色指标只表示系统是否检出双角色。",
            "official_content_bigram_* 衡量 ASR 文本与脱敏官方工单内容的字符二元组覆盖，不等价于转写准确率。",
            "category_top* 以比赛 12 类对应的规范承办部门为标签。",
            "official_handling_unit_topk 只做名称别名匹配；区县政府、中心和局委层级不同会被判为未命中。",
            "rules 模式全程本地运行，不调用外部 LLM；production 模式必须显式授权。",
        ],
    }
    total = len(aligned)
    cases: list[dict[str, Any]] = []
    if args.resume and args.output.exists():
        previous = json.loads(args.output.read_text(encoding="utf-8"))
        if previous.get("metadata", {}).get("intake_mode") != args.intake_mode:
            raise ValueError("已有报告的 intake_mode 与本次不一致，不能续跑")
        cases = list(previous.get("cases", []))
        assert_public_payload(cases)
    completed_ids = {case["id"] for case in cases}
    pending = [item for item in aligned if item["label"]["id"] not in completed_ids]
    if cases:
        print(f"resume: 已有 {len(cases)} 条，剩余 {len(pending)} 条", flush=True)
    for index, item in enumerate(pending, start=len(cases) + 1):
        print(f"[{index}/{total}] {item['label']['id']} {item['category']}", flush=True)
        result = await evaluate_case(item, asr, intake, router, args.intake_mode)
        cases.append(result)
        save_report(args.output.resolve(), metadata, cases, total)
        if result["status"] == "ok":
            print(
                f"  speech_span={result['recognized_speech_span_seconds']:.1f}s "
                f"roles={int(result['metrics']['both_roles_present'])} "
                f"fields={result['metrics']['field_accuracy']:.2f} "
                f"route={int(result['metrics']['category_routing_top1'])}",
                flush=True,
            )
        else:
            print(f"  ERROR {result['error_code']}", flush=True)
    summary = build_summary(cases, total)
    save_report(args.output.resolve(), metadata, cases, total)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset", type=Path, default=DEFAULT_DATASET)
    parser.add_argument("--labels", type=Path, default=DEFAULT_LABELS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--intake-mode", choices=("rules", "production"), default="rules")
    parser.add_argument("--allow-external-llm", action="store_true")
    parser.add_argument("--model-path", type=Path)
    parser.add_argument("--device", choices=("auto", "cuda", "cpu"), default="auto")
    parser.add_argument("--compute-type", default="auto")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--resume", action="store_true", help="从已有同模式公开报告继续未完成样例")
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    print(json.dumps(asyncio.run(run(args)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
